import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

df = pd.read_csv("data/superstore_clean.csv", dtype={"Postal_Code": str})
regions = ["Central", "East", "South", "West"]
years = [2014, 2015, 2016, 2017]
cats = ["Furniture", "Office Supplies", "Technology"]
F = "Arial"
base = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
hdr_font = Font(name=F, size=10, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F3864")
inp = Font(name=F, size=10, color="0000FF")
yellow = PatternFill("solid", fgColor="FFFF00")
note = Font(name=F, size=9, italic=True, color="595959")
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)

def hdr(ws, row, c1, labels):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=c1 + i, value=t)
        c.font, c.fill, c.border = hdr_font, hdr_fill, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

# ---------- Data sheet ----------
wd = wb.active; wd.title = "Clean_Data"
cols = list(df.columns)
wd.append(cols)
for r in df.itertuples(index=False): wd.append(list(r))
hdr(wd, 1, 1, cols)
for c in range(1, len(cols) + 1): wd.column_dimensions[L(c)].width = 15
wd.freeze_panes = "A2"
N = len(df) + 1
col = {c: L(i + 1) for i, c in enumerate(cols)}
def rng(c): return f"Clean_Data!${col[c]}$2:${col[c]}${N}"
R, Y, Q, S, P, CAT, YQ = rng("Region"), rng("Year"), rng("Quarter"), rng("Sales"), rng("Profit"), rng("Category"), rng("Year_Quarter")

# ---------- Growth table ----------
g = wb.create_sheet("Growth_Table")
g["A1"] = "Regional Sales Growth - Year over Year"; g["A1"].font = Font(name=F, size=14, bold=True)
g["A2"] = "WARNING: Year is derived from a SIMULATED order date (seed 42). The source CSV has no date column. Growth below is illustrative, not real business growth."
g["A2"].font = Font(name=F, size=10, bold=True, color="C00000")
hdr(g, 4, 1, ["Region"] + [f"Sales {y}" for y in years] + ["YoY 2015", "YoY 2016", "YoY 2017", "CAGR 2014-17", "Total Sales", "Margin %"])
for i, reg in enumerate(regions):
    r = 5 + i
    g.cell(r, 1, reg)
    for j, y in enumerate(years):
        g.cell(r, 2 + j, f'=SUMIFS({S},{R},$A{r},{Y},{y})').number_format = '#,##0'
    for j in range(3):
        cur, prv = L(3 + j), L(2 + j)
        g.cell(r, 6 + j, f'=IFERROR({cur}{r}/{prv}{r}-1,"n/a")').number_format = '0.0%;[Red]-0.0%'
    g.cell(r, 9, f'=IFERROR((E{r}/B{r})^(1/3)-1,"n/a")').number_format = '0.0%;[Red]-0.0%'
    g.cell(r, 10, f'=SUM(B{r}:E{r})').number_format = '#,##0'
    g.cell(r, 11, f'=SUMIFS({P},{R},$A{r})/J{r}').number_format = '0.0%'
t = 9
g.cell(t, 1, "Total")
for j in range(4): g.cell(t, 2 + j, f'=SUM({L(2+j)}5:{L(2+j)}8)').number_format = '#,##0'
for j in range(3):
    cur, prv = L(3 + j), L(2 + j)
    g.cell(t, 6 + j, f'=IFERROR({cur}{t}/{prv}{t}-1,"n/a")').number_format = '0.0%;[Red]-0.0%'
g.cell(t, 9, f'=IFERROR((E{t}/B{t})^(1/3)-1,"n/a")').number_format = '0.0%;[Red]-0.0%'
g.cell(t, 10, f'=SUM(B{t}:E{t})').number_format = '#,##0'
g.cell(t, 11, f'=SUM({P})/J{t}').number_format = '0.0%'
for r in range(5, 10):
    for c in range(1, 12):
        g.cell(r, c).font = bold if r == 9 else base; g.cell(r, c).border = box
g["A11"] = "Growth = (Current - Prior) / Prior. CAGR = (Sales 2017 / Sales 2014)^(1/3) - 1. Years are full calendar years (consistent periods)."
g["A11"].font = note
g.column_dimensions["A"].width = 14
for c in range(2, 12): g.column_dimensions[L(c)].width = 14
g.row_dimensions[4].height = 30

# Chart 1: sales by year (lines) ; Chart 2: YoY bars
ch = LineChart(); ch.title = "Regional Sales by Year (simulated dates)"; ch.y_axis.title = "Sales ($)"; ch.x_axis.title = "Year"
ch.height, ch.width = 9, 17
data = Reference(g, min_col=1, max_col=5, min_row=5, max_row=8)
ch.add_data(data, from_rows=True, titles_from_data=True)
ch.set_categories(Reference(g, min_col=2, max_col=5, min_row=4))
ch.y_axis.number_format = '#,##0'; ch.y_axis.delete = False; ch.x_axis.delete = False
g.add_chart(ch, "A13")
bc = BarChart(); bc.type = "col"; bc.title = "YoY Sales Growth by Region (simulated dates)"; bc.y_axis.title = "YoY growth"
bc.height, bc.width = 9, 17
bc.add_data(Reference(g, min_col=5 + 1, max_col=8, min_row=4, max_row=8), titles_from_data=True) if False else None
for j, y in enumerate([2015, 2016, 2017]):
    bc.add_data(Reference(g, min_col=6 + j, min_row=4, max_row=8), titles_from_data=True)
bc.set_categories(Reference(g, min_col=1, min_row=5, max_row=8))
bc.y_axis.number_format = '0%'; bc.y_axis.delete = False; bc.x_axis.delete = False
g.add_chart(bc, "H13")

# ---------- QoQ ----------
qs = [f"{y}-Q{q}" for y in years for q in range(1, 5)]
qq = wb.create_sheet("QoQ_Growth")
qq["A1"] = "Quarter-over-Quarter Sales Growth by Region (SIMULATED dates)"; qq["A1"].font = Font(name=F, size=14, bold=True)
hdr(qq, 3, 1, ["Year-Quarter"] + [f"{r} Sales" for r in regions] + [f"{r} QoQ" for r in regions])
for i, q in enumerate(qs):
    r = 4 + i
    qq.cell(r, 1, q)
    for j, reg in enumerate(regions):
        qq.cell(r, 2 + j, f'=SUMIFS({S},{R},"{reg}",{YQ},$A{r})').number_format = '#,##0'
        c = qq.cell(r, 6 + j, "n/a" if i == 0 else f'=IFERROR({L(2+j)}{r}/{L(2+j)}{r-1}-1,"n/a")')
        c.number_format = '0.0%;[Red]-0.0%'
    for c in range(1, 10): qq.cell(r, c).font = base; qq.cell(r, c).border = box
qq.column_dimensions["A"].width = 14
for c in range(2, 10): qq.column_dimensions[L(c)].width = 14
qq.row_dimensions[3].height = 30
lc = LineChart(); lc.title = "Quarterly Sales by Region (simulated dates)"; lc.height, lc.width = 9, 20
lc.add_data(Reference(qq, min_col=2, max_col=5, min_row=3, max_row=3 + len(qs)), titles_from_data=True)
lc.set_categories(Reference(qq, min_col=1, min_row=4, max_row=3 + len(qs)))
lc.y_axis.number_format = '#,##0'; lc.y_axis.delete = False; lc.x_axis.delete = False
qq.add_chart(lc, "K3")

# ---------- Small base ----------
sb = wb.create_sheet("Small_Base_Check")
sb["A1"] = "Small-Base Check: Region x Category YoY growth (SIMULATED dates)"; sb["A1"].font = Font(name=F, size=14, bold=True)
sb["A2"] = "Small-base threshold ($) - edit:"; sb["A2"].font = bold
sb["D2"] = 40000; sb["D2"].font = inp; sb["D2"].fill = yellow; sb["D2"].number_format = '#,##0'
sb["E2"] = "ASSUMPTION: prior-year sales below this value are flagged as a small base. Chosen as roughly the lower quartile of region-category-year sales."
sb["E2"].font = note
hdr(sb, 4, 1, ["Region", "Category"] + [f"Sales {y}" for y in years] + ["YoY 2015", "YoY 2016", "YoY 2017", "Smallest prior base", "Small-base flag"])
r = 5
for reg in regions:
    for cat in cats:
        sb.cell(r, 1, reg); sb.cell(r, 2, cat)
        for j, y in enumerate(years):
            sb.cell(r, 3 + j, f'=SUMIFS({S},{R},$A{r},{CAT},$B{r},{Y},{y})').number_format = '#,##0'
        for j in range(3):
            cur, prv = L(4 + j), L(3 + j)
            sb.cell(r, 7 + j, f'=IFERROR({cur}{r}/{prv}{r}-1,"n/a")').number_format = '0.0%;[Red]-0.0%'
        sb.cell(r, 10, f'=MIN(C{r}:E{r})').number_format = '#,##0'
        sb.cell(r, 11, f'=IF(J{r}<$D$2,"SMALL BASE","ok")')
        for c in range(1, 12): sb.cell(r, c).font = base; sb.cell(r, c).border = box
        r += 1
sb.conditional_formatting.add(f"K5:K{r-1}", CellIsRule(operator="equal", formula=['"SMALL BASE"'], fill=PatternFill("solid", bgColor="FFC7CE")))
sb.column_dimensions["A"].width = 14; sb.column_dimensions["B"].width = 18
for c in range(3, 12): sb.column_dimensions[L(c)].width = 14
sb.row_dimensions[4].height = 30

# ---------- Cross-section (real) ----------
cs = wb.create_sheet("Region_Profile_REAL")
cs["A1"] = "Region profile - uses NO dates, so these figures are real"; cs["A1"].font = Font(name=F, size=14, bold=True)
hdr(cs, 3, 1, ["Region", "Sales", "Sales share", "Profit", "Margin %", "Avg discount", "Loss-making lines %"])
D, LM = rng("Discount"), rng("Loss_Making")
for i, reg in enumerate(regions):
    r = 4 + i
    cs.cell(r, 1, reg)
    cs.cell(r, 2, f'=SUMIFS({S},{R},$A{r})').number_format = '#,##0'
    cs.cell(r, 3, f'=B{r}/SUM($B$4:$B$7)').number_format = '0.0%'
    cs.cell(r, 4, f'=SUMIFS({P},{R},$A{r})').number_format = '#,##0'
    cs.cell(r, 5, f'=D{r}/B{r}').number_format = '0.0%'
    cs.cell(r, 6, f'=AVERAGEIFS({D},{R},$A{r})').number_format = '0.0%'
    cs.cell(r, 7, f'=AVERAGEIFS({LM},{R},$A{r})').number_format = '0.0%'
    for c in range(1, 8): cs.cell(r, c).font = base; cs.cell(r, c).border = box
cs.column_dimensions["A"].width = 14
for c in range(2, 8): cs.column_dimensions[L(c)].width = 16

# ---------- Notes ----------
n = wb.create_sheet("Notes_Assumptions", 0)
lines = [
 ("Task 30 - Regional Growth Analysis: workbook notes", True),
 ("", False),
 ("DATA", True),
 ("Source: SampleSuperstore.csv (9,994 rows, 13 columns). After removing 17 exact duplicates: 9,977 rows (sheet Clean_Data).", False),
 ("The file has NO order-date column.", False),
 ("ASSUMPTION (hardcoded): Order_Date_SIMULATED was assigned randomly (numpy seed 42) between 2014-01-01 and 2017-12-31.", False),
 ("Consequence: growth rates on Growth_Table, QoQ_Growth and Small_Base_Check show sampling noise, not true business trends.", False),
 ("Region_Profile_REAL uses no dates, so it is genuine.", False),
 ("To use real dates: replace the Order_Date_SIMULATED/Year/Quarter/Year_Quarter columns in Clean_Data; all formulas recalculate.", False),
 ("", False),
 ("DEFINITIONS", True),
 ("YoY growth = (Sales this year - Sales prior year) / Sales prior year", False),
 ("CAGR = (Sales 2017 / Sales 2014)^(1/3) - 1", False),
 ("Small base = prior-period sales below the threshold in Small_Base_Check!D2 (blue/yellow input cell).", False),
 ("Colour code: blue text = hardcoded input; black = formula.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = n.cell(i, 1, t); c.font = Font(name=F, size=12 if i == 1 else 10, bold=b)
n.column_dimensions["A"].width = 130
wb.save("outputs/Task30_Regional_Growth_Analysis.xlsx")
