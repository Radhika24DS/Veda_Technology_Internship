"""Step 2 - Build the Excel deliverable (formula-driven) from the cleaned data."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule

df = pd.read_csv("data/superstore_clean.csv", dtype={"Postal Code": str})
log = pd.read_csv("data/cleaning_log.csv")
N = len(df); LAST = N + 1

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF", size=10)
HF = PatternFill("solid", fgColor="1F3A5F")
B = Font(name=F, size=10); BB = Font(name=F, size=10, bold=True)
T = Font(name=F, size=14, bold=True, color="1F3A5F")
INP = Font(name=F, size=10, color="0000FF"); YEL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col+i, t); c.font = H; c.fill = HF
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    ws.row_dimensions[row].height = 42
def setw(ws, widths):
    for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
def style(ws, r1, r2, c1, c2, fmts=None):
    for r in range(r1, r2+1):
        for c in range(c1, c2+1):
            x = ws.cell(r, c); x.font = B; x.border = BOX
            if fmts and c in fmts: x.number_format = fmts[c]

wb = Workbook()

# ---------------- About ----------------
ws = wb.active; ws.title = "About"
rows = [
 ("Task 27 - Salesperson Performance Analysis (Superstore)", T),
 ("Prepared for: Veda Technology Internship - Data Analytics Track", B),
 ("", B),
 ("IMPORTANT DATA LIMITATIONS (read first)", BB),
 ("1. The dataset has NO salesperson column. 'Salesperson' is modelled two ways: Region Team (4 units) and State Rep (49 units). Replace with a real mapping table if one exists.", B),
 ("2. The dataset has NO order date, so period-over-period GROWTH cannot be calculated. Growth is carried in the scorecard with weight 0% and is marked N/A.", B),
 ("3. The dataset has NO Order ID, so true Average Order Value is not possible. 'AOV (proxy)' = Sales per line item.", B),
 ("4. The dataset has no territory size data (population, accounts, quota). Territory differences are handled with a product-mix benchmark and size tiers.", B),
 ("", B),
 ("SHEETS", BB),
 ("Weights - editable scorecard weights (blue cells). Everything recalculates.", B),
 ("Region_Ranking - 4 Region Teams: KPIs, normalised scores, composite score, rank, sales-only rank vs composite rank.", B),
 ("State_Ranking - 49 State Reps: same KPIs plus size tier and rank within tier.", B),
 ("Leaderboard - sorted ranking views driven by formulas.", B),
 ("Region_vs_Category - margin by region x category; discount-band mix by region.", B),
 ("SubCat_Benchmark - national margin per sub-category used as the 'expected margin' benchmark.", B),
 ("Clean_Data - 9,977 cleaned rows (Expected Profit column is a live formula).", B),
 ("Cleaning_Log - every preprocessing step and rows affected.", B),
 ("Recommendations - findings and actions.", B),
 ("", B),
 ("METHOD", BB),
 ("Composite score = 100 x sum(weight x min-max normalised KPI). Average Discount is inverted (lower = better). Margin Gap = actual margin minus margin expected if the unit had sold its product mix at national sub-category margins (removes product-mix bias).", B),
]
for i, (t, f) in enumerate(rows, 1):
    c = ws.cell(i, 1, t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 130

# ---------------- Weights ----------------
wsw = wb.create_sheet("Weights")
wsw["A1"] = "Scorecard Weights"; wsw["A1"].font = T
wsw["A2"] = "Blue cells on yellow are inputs. Weights must total 100%."; wsw["A2"].font = B
hdr(wsw, 3, ["KPI", "Weight", "Why"])
W = [("Sales", .20, "Scale of business; deliberately not the biggest weight"),
     ("Profit", .25, "Value actually created"),
     ("Profit Margin", .20, "Efficiency of each sales dollar"),
     ("AOV (proxy: sales per line item)", .10, "Deal size / selling-up ability"),
     ("Discount discipline (inverted avg discount)", .10, "Does the rep buy sales with discounts?"),
     ("Margin Gap vs mix benchmark", .15, "Fair comparison: removes product-mix / territory advantage"),
     ("Growth (N/A - no date field)", 0.0, "Cannot be computed; set to 0%. Re-weight if dates become available")]
for i, (k, w, why) in enumerate(W, 4):
    wsw.cell(i, 1, k).font = B; c = wsw.cell(i, 2, w); c.font = INP; c.fill = YEL; c.number_format = "0%"
    wsw.cell(i, 3, why).font = B
wsw["A11"] = "Total (must be 100%)"; wsw["A11"].font = BB
wsw["B11"] = "=SUM(B4:B10)"; wsw["B11"].number_format = "0%"; wsw["B11"].font = BB
wsw["C11"] = '=IF(ROUND(B11,4)=1,"OK","WEIGHTS DO NOT SUM TO 100%")'; wsw["C11"].font = BB
wsw["A13"] = "State size-tier thresholds (percentile of state sales)"; wsw["A13"].font = BB
wsw["A14"] = "Small / Mid cut-off percentile"; wsw["B14"] = 0.33
wsw["A15"] = "Mid / Large cut-off percentile"; wsw["B15"] = 0.67
for r in (14, 15):
    wsw.cell(r, 1).font = B; wsw.cell(r, 2).font = INP; wsw.cell(r, 2).fill = YEL; wsw.cell(r, 2).number_format = "0%"
wsw["A16"] = "Small / Mid sales cut-off ($)"; wsw["B16"] = "=PERCENTILE(State_Ranking!$C$5:$C$53,B14)"
wsw["A17"] = "Mid / Large sales cut-off ($)"; wsw["B17"] = "=PERCENTILE(State_Ranking!$C$5:$C$53,B15)"
wsw["A18"] = "Low-volume warning below (line items)"; wsw["B18"] = 50
wsw["B18"].font = INP; wsw["B18"].fill = YEL
for r in (16, 17): wsw.cell(r, 1).font = B; wsw.cell(r, 2).font = B; wsw.cell(r, 2).number_format = "$#,##0"
wsw["A18"].font = B
setw(wsw, [46, 14, 70])

# ---------------- Clean_Data ----------------
wsd = wb.create_sheet("Clean_Data")
cols = ["Ship Mode","Segment","City","State","Postal Code","Region","Category","Sub-Category","Sales","Quantity",
        "Discount","Profit","Profit Margin","Loss Flag","Discount Band","Salesperson (Region Team)","Salesperson (State Rep)","Profit Outlier"]
hdr(wsd, 1, cols + ["Expected Profit (mix benchmark)"])
for r, row in enumerate(df[cols].itertuples(index=False), 2):
    for c, v in enumerate(row, 1):
        wsd.cell(r, c, v)
    wsd.cell(r, 19, f"=I{r}*INDEX(SubCat_Benchmark!$F$5:$F$21,MATCH(H{r},SubCat_Benchmark!$A$5:$A$21,0))")
for c, w in enumerate([14,12,18,18,11,10,16,14,11,9,10,11,10,8,13,24,28,9,16], 1):
    wsd.column_dimensions[L(c)].width = w
wsd.freeze_panes = "A2"; wsd.auto_filter.ref = f"A1:S{LAST}"
for col, fmt in {"I": "#,##0.00", "L": "#,##0.00", "K": "0%", "M": "0.0%", "S": "#,##0.00"}.items():
    for r in range(2, LAST+1): wsd[f"{col}{r}"].number_format = fmt
rng = lambda col: f"Clean_Data!${col}$2:${col}${LAST}"

# ---------------- SubCat_Benchmark ----------------
wsb = wb.create_sheet("SubCat_Benchmark")
wsb["A1"] = "National margin by sub-category (benchmark for fair comparison)"; wsb["A1"].font = T
wsb["A2"] = "Expected margin = what a unit would earn if it sold each sub-category at the national margin."; wsb["A2"].font = B
hdr(wsb, 4, ["Sub-Category", "Category", "Sales", "Profit", "Line items", "National Margin"])
subs = df.groupby(["Sub-Category", "Category"]).size().reset_index()[["Sub-Category", "Category"]]
for i, (s, c) in enumerate(subs.itertuples(index=False), 5):
    wsb.cell(i, 1, s); wsb.cell(i, 2, c)
    wsb.cell(i, 3, f"=SUMIFS({rng('I')},{rng('H')},A{i})")
    wsb.cell(i, 4, f"=SUMIFS({rng('L')},{rng('H')},A{i})")
    wsb.cell(i, 5, f"=COUNTIFS({rng('H')},A{i})")
    wsb.cell(i, 6, f"=D{i}/C{i}")
assert len(subs) == 17
style(wsb, 5, 21, 1, 6, {3: "$#,##0", 4: "$#,##0", 5: "#,##0", 6: "0.0%"})
wsb.conditional_formatting.add("F5:F21", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
setw(wsb, [16, 18, 14, 14, 12, 16])

# ---------------- ranking builder ----------------
def build_rank(ws, title, keycol, labelcol, keys, first, with_tier):
    ws["A1"] = title; ws["A1"].font = T
    ws["A2"] = "All figures are live formulas on Clean_Data. Growth is N/A (no date field). AOV is a proxy (no Order ID)."; ws["A2"].font = B
    heads = ["Salesperson", "Key" if not with_tier else "Region", "Sales", "Profit", "Profit Margin", "Line Items",
             "AOV (proxy)", "Avg Discount", "% Loss-Making Lines", "Expected Profit", "Expected Margin", "Margin Gap (pp)",
             "n: Sales", "n: Profit", "n: Margin", "n: AOV", "n: Discount (inv)", "n: Margin Gap",
             "Composite Score", "Rank (Composite)", "Rank (Sales only)", "Rank change vs sales-only"]
    if with_tier: heads += ["Size Tier", "Rank within Tier", "Data Quality"]
    hdr(ws, 4, heads)
    last = first + len(keys) - 1
    for i, (lab, key, extra) in enumerate(keys, first):
        ws.cell(i, 1, lab)
        ws.cell(i, 2, key if not with_tier else extra)
        crit = f"{rng(labelcol)},$A{i}"
        ws.cell(i, 3, f"=SUMIFS({rng('I')},{crit})")
        ws.cell(i, 4, f"=SUMIFS({rng('L')},{crit})")
        ws.cell(i, 5, f"=D{i}/C{i}")
        ws.cell(i, 6, f"=COUNTIFS({crit})")
        ws.cell(i, 7, f"=C{i}/F{i}")
        ws.cell(i, 8, f"=AVERAGEIFS({rng('K')},{crit})")
        ws.cell(i, 9, f"=AVERAGEIFS({rng('N')},{crit})")
        ws.cell(i, 10, f"=SUMIFS({rng('S')},{crit})")
        ws.cell(i, 11, f"=J{i}/C{i}")
        ws.cell(i, 12, f"=E{i}-K{i}")
        for col, src in zip("MNOPQR", "CDEGHL"):
            rg = f"{src}${first}:{src}${last}"
            base = f"({src}{i}-MIN({rg}))/(MAX({rg})-MIN({rg}))"
            ws[f"{col}{i}"] = f"=1-{base}" if src == "H" else f"={base}"
        ws.cell(i, 19, f"=100*(M{i}*Weights!$B$4+N{i}*Weights!$B$5+O{i}*Weights!$B$6+P{i}*Weights!$B$7+Q{i}*Weights!$B$8+R{i}*Weights!$B$9)")
        ws.cell(i, 20, f"=RANK(S{i},S${first}:S${last})")
        ws.cell(i, 21, f"=RANK(C{i},C${first}:C${last})")
        ws.cell(i, 22, f"=U{i}-T{i}")
        if with_tier:
            ws.cell(i, 23, f'=IF(C{i}<Weights!$B$16,"Small",IF(C{i}<Weights!$B$17,"Mid","Large"))')
            ws.cell(i, 24, f'=COUNTIFS(W${first}:W${last},W{i},S${first}:S${last},">"&S{i})+1')
            ws.cell(i, 25, f'=IF(F{i}<Weights!$B$18,"Low volume - interpret with caution","OK")')
    tot = last + 1
    ws.cell(tot, 1, "Company total / average (reference)").font = BB
    ws.cell(tot, 3, f"=SUM(C{first}:C{last})"); ws.cell(tot, 4, f"=SUM(D{first}:D{last})")
    ws.cell(tot, 5, f"=D{tot}/C{tot}"); ws.cell(tot, 6, f"=SUM(F{first}:F{last})"); ws.cell(tot, 7, f"=C{tot}/F{tot}")
    ws.cell(tot, 8, f"=AVERAGE(Clean_Data!$K$2:$K${LAST})"); ws.cell(tot, 9, f"=AVERAGE(Clean_Data!$N$2:$N${LAST})")
    ws.cell(tot, 10, f"=SUM(J{first}:J{last})"); ws.cell(tot, 11, f"=J{tot}/C{tot}"); ws.cell(tot, 12, f"=E{tot}-K{tot}")
    fm = {3:"$#,##0",4:"$#,##0",5:"0.0%",6:"#,##0",7:"$#,##0.00",8:"0.0%",9:"0.0%",10:"$#,##0",11:"0.0%",12:"0.0%",
          13:"0.00",14:"0.00",15:"0.00",16:"0.00",17:"0.00",18:"0.00",19:"0.0",22:"+0;-0;0"}
    style(ws, first, tot, 1, len(heads), fm)
    for c in range(1, len(heads)+1): ws.cell(tot, c).font = BB
    ws.conditional_formatting.add(f"S{first}:S{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="63BE7B"))
    ws.conditional_formatting.add(f"E{first}:E{last}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0.05, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    ws.freeze_panes = ws.cell(first, 3)
    return last

# Region Ranking
regions = sorted(df.Region.unique())
build_rank(wb.create_sheet("Region_Ranking"), "Region Team Ranking (4 salespeople proxies)", "Region", "P",
           [(f"{r} Region Team", r, None) for r in regions], 5, False)
setw(wb["Region_Ranking"], [26, 10] + [13]*20)

# State ranking (sorted by sales desc for readability)
st = df.groupby("State").agg(S=("Sales", "sum"), R=("Region", "first")).sort_values("S", ascending=False)
assert len(st) == 49 and df.groupby("State").Region.nunique().max() == 1
build_rank(wb.create_sheet("State_Ranking"), "State Rep Ranking (49 salespeople proxies)", "State", "Q",
           [(f"State Rep - {s}", s, r) for s, r in zip(st.index, st.R)], 5, True)
setw(wb["State_Ranking"], [34, 10] + [13]*20 + [10, 12, 34])

# ---------------- Leaderboard ----------------
wl = wb.create_sheet("Leaderboard")
wl["A1"] = "Leaderboard (formula-driven, sorts itself when weights change)"; wl["A1"].font = T
wl["A3"] = "Region Teams - by composite score"; wl["A3"].font = BB
hdr(wl, 4, ["Rank", "Salesperson", "Composite Score", "Sales", "Profit", "Profit Margin", "AOV (proxy)", "Avg Discount", "Margin Gap (pp)", "Rank if judged on Sales only"])
for k in range(1, 5):
    r = 4 + k; wl.cell(r, 1, k)
    m = f"MATCH($A{r},Region_Ranking!$T$5:$T$8,0)"
    for c, src in zip(range(2, 11), "ASCDEGHLU"):
        wl.cell(r, c, f"=INDEX(Region_Ranking!${src}$5:${src}$8,{m})")
style(wl, 5, 8, 1, 10, {3: "0.0", 4: "$#,##0", 5: "$#,##0", 6: "0.0%", 7: "$#,##0.00", 8: "0.0%", 9: "0.0%"})
wl["A11"] = "State Reps - Top 10 by composite score"; wl["A11"].font = BB
hdr(wl, 12, ["Rank", "Salesperson", "Composite Score", "Sales", "Profit", "Profit Margin", "AOV (proxy)", "Avg Discount", "Margin Gap (pp)", "Size Tier / Data Quality"])
for k in range(1, 11):
    r = 12 + k; wl.cell(r, 1, k)
    m = f"MATCH($A{r},State_Ranking!$T$5:$T$53,0)"
    for c, src in zip(range(2, 10), "ASCDEGHL"):
        wl.cell(r, c, f"=INDEX(State_Ranking!${src}$5:${src}$53,{m})")
    wl.cell(r, 10, f'=INDEX(State_Ranking!$W$5:$W$53,{m})&" / "&INDEX(State_Ranking!$Y$5:$Y$53,{m})')
style(wl, 13, 22, 1, 10, {3: "0.0", 4: "$#,##0", 5: "$#,##0", 6: "0.0%", 7: "$#,##0.00", 8: "0.0%", 9: "0.0%"})
wl["A25"] = "State Reps - Bottom 10 by composite score"; wl["A25"].font = BB
hdr(wl, 26, ["Rank", "Salesperson", "Composite Score", "Sales", "Profit", "Profit Margin", "AOV (proxy)", "Avg Discount", "Margin Gap (pp)", "Size Tier / Data Quality"])
for k in range(1, 11):
    r = 26 + k; wl.cell(r, 1, f"=COUNT(State_Ranking!$T$5:$T$53)-{10-k}")
    m = f"MATCH($A{r},State_Ranking!$T$5:$T$53,0)"
    for c, src in zip(range(2, 10), "ASCDEGHL"):
        wl.cell(r, c, f"=INDEX(State_Ranking!${src}$5:${src}$53,{m})")
    wl.cell(r, 10, f'=INDEX(State_Ranking!$W$5:$W$53,{m})&" / "&INDEX(State_Ranking!$Y$5:$Y$53,{m})')
style(wl, 27, 36, 1, 10, {3: "0.0", 4: "$#,##0", 5: "$#,##0", 6: "0.0%", 7: "$#,##0.00", 8: "0.0%", 9: "0.0%"})
setw(wl, [8, 32, 14, 14, 14, 13, 13, 13, 14, 38])

# ---------------- Region_vs_Category ----------------
wr = wb.create_sheet("Region_vs_Category")
wr["A1"] = "Territory differences: where each region makes and loses money"; wr["A1"].font = T
wr["A3"] = "Profit margin by Region x Category"; wr["A3"].font = BB
cats = sorted(df.Category.unique())
hdr(wr, 4, ["Region"] + cats + ["All categories"])
for i, r in enumerate(regions, 5):
    wr.cell(i, 1, r)
    for j, c in enumerate(cats, 2):
        wr.cell(i, j, f"=SUMIFS({rng('L')},{rng('F')},$A{i},{rng('G')},{L(j)}$4)/SUMIFS({rng('I')},{rng('F')},$A{i},{rng('G')},{L(j)}$4)")
    wr.cell(i, 5, f"=SUMIFS({rng('L')},{rng('F')},$A{i})/SUMIFS({rng('I')},{rng('F')},$A{i})")
style(wr, 5, 8, 1, 5, {2: "0.0%", 3: "0.0%", 4: "0.0%", 5: "0.0%"})
wr.conditional_formatting.add("B5:E8", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
wr["A11"] = "Share of line items by discount band (per region)"; wr["A11"].font = BB
bands = ["No discount", "1-10%", "11-20%", "21-30%", ">30%"]
hdr(wr, 12, ["Region"] + bands + ["Profit on >30% discount lines ($)"])
for i, r in enumerate(regions, 13):
    wr.cell(i, 1, r)
    for j, b in enumerate(bands, 2):
        wr.cell(i, j, '=COUNTIFS(' + rng('F') + ',$A' + str(i) + ',' + rng('O') + ',"="&' + L(j) + '$12)/COUNTIFS(' + rng('F') + ',$A' + str(i) + ')')
    wr.cell(i, 7, f'=SUMIFS({rng("L")},{rng("F")},$A{i},{rng("O")},"=>30%")')
style(wr, 13, 16, 1, 7, {2: "0.0%", 3: "0.0%", 4: "0.0%", 5: "0.0%", 6: "0.0%", 7: "$#,##0;($#,##0)"})
wr.conditional_formatting.add("F13:F16", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F8696B"))
wr["A19"] = "Profit ($) by Region x Sub-Category"; wr["A19"].font = BB
hdr(wr, 20, ["Sub-Category"] + regions + ["Total"])
for i, s in enumerate(subs["Sub-Category"], 21):
    wr.cell(i, 1, s)
    for j, r in enumerate(regions, 2):
        wr.cell(i, j, f"=SUMIFS({rng('L')},{rng('H')},$A{i},{rng('F')},{L(j)}$20)")
    wr.cell(i, 6, f"=SUM(B{i}:E{i})")
style(wr, 21, 37, 1, 6, {k: "$#,##0;($#,##0)" for k in range(2, 7)})
wr.conditional_formatting.add("B21:E37", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
setw(wr, [20, 14, 14, 14, 14, 14, 24])

# ---------------- Cleaning_Log ----------------
wc = wb.create_sheet("Cleaning_Log")
wc["A1"] = "Data Cleaning & Preprocessing Log"; wc["A1"].font = T
hdr(wc, 3, ["Step", "Detail", "Rows affected"])
for i, row in enumerate(log.itertuples(index=False), 4):
    for j, v in enumerate(row, 1): wc.cell(i, j, v)
style(wc, 4, 3 + len(log), 1, 3, {3: "#,##0"}); setw(wc, [24, 95, 14])


# ---------------- Recommendations ----------------
wrec = wb.create_sheet("Recommendations")
wrec["A1"] = "Findings and Recommendations (figures from the cleaned dataset, 9,977 rows)"; wrec["A1"].font = T
hdr(wrec, 3, ["#", "Finding (evidence)", "Recommendation", "Owner / Priority"])
R = [
 ("West ranks #1 overall: highest sales ($725K), profit ($108K), margin (14.9%), lowest avg discount (11%) and only 3.7% of lines discounted over 30%.",
  "Document West's discount and deal-approval practices and use them as the playbook for other regions.", "Sales Ops / High"),
 ("Central ranks last on composite score (6.5/100) even though its sales rank is 3rd. Margin is 7.9%, avg discount 24%, 21% of lines are discounted over 30% and 32% of lines lose money. Its product mix would predict a 13.3% margin (highest of the four), so mix does not explain the shortfall.",
  "Introduce a discount cap and manager approval above 20% in Central; review the 495 lines with >30% discount.", "Regional Head / High"),
 ("Discounts above 30% generate -$125K on $260K sales across all regions, while undiscounted lines earn +$321K.",
  "Company-wide approval gate for >30% discounts; track Discount Band mix per rep monthly.", "Finance + Sales / High"),
 ("Revenue alone misranks reps: South is 4th on sales but 3rd on composite; Central is 3rd on sales but last on composite.",
  "Use the composite scorecard (not revenue) for incentives and reviews; keep weights visible and agreed.", "Sales Leadership / Medium"),
 ("Furniture loses money in Central (-1.8% margin) and is weak everywhere; Tables is the worst sub-category (about -$17.7K), mostly in East, Central and South.",
  "Review Table pricing and Furniture discounting.", "Category Mgmt / Medium"),
 ("State-level: Ohio, Colorado, Illinois, Tennessee and Texas rank lowest; Texas alone lost $25.8K on $170K sales. California and New York lead.",
  "Targeted deal reviews in the bottom-10 states; pair mid-ranked reps with California / New York reps for coaching.", "Regional Heads / Medium"),
 ("Territory fairness: ranking uses a product-mix margin benchmark and size tiers, so a rep in a Furniture-heavy territory is not unfairly penalised.",
  "Add real territory data (accounts, population, quota) and an Order Date column to enable growth and quota-attainment KPIs.", "BI Team / Medium"),
 ("Data gaps: no Salesperson, Order Date or Order ID in the file.",
  "Request a rep-to-territory mapping and order-level data; then replace Region/State proxies and add true AOV and YoY growth.", "Data Owner / High"),
]
for i, (a, b, c) in enumerate(R, 4):
    for j, v in enumerate([i-3, a, b, c], 1):
        x = wrec.cell(i, j, v); x.font = B; x.border = BOX; x.alignment = Alignment(wrap_text=True, vertical="top")
    wrec.row_dimensions[i].height = 62
setw(wrec, [5, 80, 70, 24])

wb.save("excel/Salesperson_Performance_Analysis.xlsx")
print("saved")
